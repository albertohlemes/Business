"""
Gerenciador de Templates para Minutas - Versão Fiel ao Template
Preserva cabeçalho, rodapé, fontes e espaçamentos do documento modelo
"""
import os
import re
import io
import copy
from datetime import datetime
from typing import Dict, Any, Optional, List
from docx import Document
from docx.shared import Pt, Inches, Cm, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import logging

logger = logging.getLogger(__name__)

TEMPLATES_DIR = "/app/backend/uploads/templates"
os.makedirs(TEMPLATES_DIR, exist_ok=True)


class TemplateManagerFiel:
    """Gerencia templates preservando formatação fiel ao original"""
    
    def __init__(self):
        self.templates_dir = TEMPLATES_DIR
    
    def extrair_formatacao_completa(self, template_path: str) -> Dict[str, Any]:
        """Extrai TODA a formatação do template"""
        formato = {
            "fonte": {
                "nome": "Times New Roman",
                "tamanho": 12,
                "cor": None
            },
            "margens": {
                "top": 2.5,
                "bottom": 2.5,
                "left": 3.0,
                "right": 2.0
            },
            "espacamento": {
                "linha": 1.5,
                "antes_paragrafo": 0,
                "depois_paragrafo": 6
            },
            "cabecalho": None,
            "rodape": None,
            "estilos_paragrafo": {}
        }
        
        if not os.path.exists(template_path):
            return formato
        
        try:
            doc = Document(template_path)
            
            # Extrair margens da seção
            section = doc.sections[0]
            if section.top_margin:
                formato["margens"]["top"] = section.top_margin.cm
            if section.bottom_margin:
                formato["margens"]["bottom"] = section.bottom_margin.cm
            if section.left_margin:
                formato["margens"]["left"] = section.left_margin.cm
            if section.right_margin:
                formato["margens"]["right"] = section.right_margin.cm
            
            # Extrair fonte padrão do primeiro parágrafo com conteúdo
            for para in doc.paragraphs:
                if para.text.strip() and para.runs:
                    run = para.runs[0]
                    if run.font.name:
                        formato["fonte"]["nome"] = run.font.name
                    if run.font.size:
                        formato["fonte"]["tamanho"] = run.font.size.pt
                    
                    # Espaçamento de linha
                    if para.paragraph_format.line_spacing:
                        formato["espacamento"]["linha"] = para.paragraph_format.line_spacing
                    if para.paragraph_format.space_before:
                        formato["espacamento"]["antes_paragrafo"] = para.paragraph_format.space_before.pt
                    if para.paragraph_format.space_after:
                        formato["espacamento"]["depois_paragrafo"] = para.paragraph_format.space_after.pt
                    break
            
            # Extrair cabeçalho
            if section.header and section.header.paragraphs:
                cabecalho_texto = []
                for para in section.header.paragraphs:
                    if para.text.strip():
                        cabecalho_texto.append(para.text)
                if cabecalho_texto:
                    formato["cabecalho"] = {
                        "texto": cabecalho_texto,
                        "alinhamento": "center"
                    }
            
            # Extrair rodapé
            if section.footer and section.footer.paragraphs:
                rodape_texto = []
                for para in section.footer.paragraphs:
                    if para.text.strip():
                        rodape_texto.append(para.text)
                if rodape_texto:
                    formato["rodape"] = {
                        "texto": rodape_texto,
                        "alinhamento": "center"
                    }
            
            logger.info(f"Formatação extraída: fonte={formato['fonte']['nome']} {formato['fonte']['tamanho']}pt")
            
        except Exception as e:
            logger.error(f"Erro ao extrair formatação: {e}")
        
        return formato
    
    def gerar_documento_fiel(
        self,
        conteudo: str,
        template_path: str,
        dados_extraidos: Dict[str, Any] = None
    ) -> bytes:
        """
        Gera documento FIEL ao template:
        - Mesmo cabeçalho e rodapé (preservados diretamente do template)
        - Mesma fonte e tamanho
        - Mesmo espaçamento
        - Mesmas margens
        """
        
        # Se tem template, usa como base
        if template_path and os.path.exists(template_path):
            # Copiar o template como base - preserva tudo (cabeçalho, rodapé, estilos)
            doc = Document(template_path)
            formato = self.extrair_formatacao_completa(template_path)
            
            # Verificar se o template tem cabeçalho/rodapé definido
            template_tem_header = bool(formato.get("cabecalho"))
            template_tem_footer = bool(formato.get("rodape"))
            
            # Limpar APENAS o conteúdo do corpo (preserva cabeçalho/rodapé nativos)
            for para in list(doc.paragraphs):
                p = para._element
                p.getparent().remove(p)
            
            # Adicionar conteúdo preservando formatação
            self._adicionar_conteudo_formatado(doc, conteudo, formato, dados_extraidos)
            
            # Só adiciona cabeçalho/rodapé se o template NÃO tinha
            # (o template original já foi preservado ao copiar o documento)
            if not template_tem_header and not template_tem_footer:
                # Se não tinha, usa o padrão Business
                formato_padrao = self._formato_padrao()
                self._garantir_cabecalho_rodape(doc, formato_padrao)
        else:
            # Criar documento novo com formatação padrão
            doc = Document()
            formato = self._formato_padrao()
            self._configurar_secao(doc, formato)
            
            # Adicionar conteúdo
            self._adicionar_conteudo_formatado(doc, conteudo, formato, dados_extraidos)
            
            # Adicionar cabeçalho e rodapé padrão
            self._garantir_cabecalho_rodape(doc, formato)
        
        # Salvar em memória
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        
        return buffer.getvalue()
    
    def _formato_padrao(self) -> Dict[str, Any]:
        """Retorna formatação padrão Business Contabilidade"""
        return {
            "fonte": {
                "nome": "Times New Roman",
                "tamanho": 12
            },
            "margens": {
                "top": 2.5,
                "bottom": 2.5,
                "left": 3.0,
                "right": 2.0
            },
            "espacamento": {
                "linha": 1.5,
                "antes_paragrafo": 0,
                "depois_paragrafo": 6
            },
            "cabecalho": {
                "texto": ["BUSINESS CONTABILIDADE", "Portal Societário"],
                "alinhamento": "center"
            },
            "rodape": {
                "texto": ["Documento gerado pelo Portal Societário"],
                "alinhamento": "center"
            }
        }
    
    def _configurar_secao(self, doc: Document, formato: Dict[str, Any]):
        """Configura seção com margens do template"""
        section = doc.sections[0]
        
        margens = formato.get("margens", {})
        section.top_margin = Cm(margens.get("top", 2.5))
        section.bottom_margin = Cm(margens.get("bottom", 2.5))
        section.left_margin = Cm(margens.get("left", 3.0))
        section.right_margin = Cm(margens.get("right", 2.0))
        
        # Configurar para cabeçalho/rodapé aparecer em todas as páginas
        section.different_first_page_header_footer = False
    
    def _garantir_cabecalho_rodape(self, doc: Document, formato: Dict[str, Any]):
        """Adiciona/atualiza cabeçalho e rodapé em todas as páginas"""
        section = doc.sections[0]
        
        # Cabeçalho
        cabecalho_config = formato.get("cabecalho")
        if cabecalho_config:
            header = section.header
            header.is_linked_to_previous = False
            
            # Limpar cabeçalho existente
            for para in header.paragraphs:
                para.clear()
            
            # Adicionar texto do cabeçalho
            textos = cabecalho_config.get("texto", [])
            for i, texto in enumerate(textos):
                if i == 0:
                    para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
                else:
                    para = header.add_paragraph()
                
                run = para.add_run(texto)
                run.font.name = formato["fonte"]["nome"]
                run.font.size = Pt(formato["fonte"]["tamanho"] - 2)  # Menor que o corpo
                run.bold = (i == 0)  # Primeiro item em negrito
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Rodapé
        rodape_config = formato.get("rodape")
        if rodape_config:
            footer = section.footer
            footer.is_linked_to_previous = False
            
            # Limpar rodapé existente
            for para in footer.paragraphs:
                para.clear()
            
            # Adicionar texto do rodapé
            textos = rodape_config.get("texto", [])
            for i, texto in enumerate(textos):
                if i == 0:
                    para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
                else:
                    para = footer.add_paragraph()
                
                run = para.add_run(texto)
                run.font.name = formato["fonte"]["nome"]
                run.font.size = Pt(9)  # Rodapé menor
                run.italic = True
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    def _adicionar_conteudo_formatado(
        self,
        doc: Document,
        conteudo: str,
        formato: Dict[str, Any],
        dados_extraidos: Dict[str, Any] = None
    ):
        """Adiciona conteúdo ao documento preservando formatação do template"""
        
        fonte_nome = formato["fonte"]["nome"]
        fonte_tamanho = formato["fonte"]["tamanho"]
        espacamento = formato.get("espacamento", {})
        
        linhas = conteudo.split('\n')
        
        for linha in linhas:
            linha_limpa = linha.strip()
            
            if not linha_limpa:
                # Linha vazia - adicionar espaço
                para = doc.add_paragraph()
                para.paragraph_format.space_after = Pt(espacamento.get("depois_paragrafo", 6))
                continue
            
            # Criar parágrafo
            para = doc.add_paragraph()
            
            # Detectar tipo de linha
            is_titulo = self._is_titulo(linha_limpa)
            is_clausula = self._is_clausula(linha_limpa)
            is_assinatura = self._is_assinatura(linha_limpa)
            
            # Adicionar texto com formatação
            run = para.add_run(linha_limpa)
            run.font.name = fonte_nome
            
            # Aplicar estilo baseado no tipo
            if is_titulo:
                run.bold = True
                run.font.size = Pt(fonte_tamanho + 2)
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.space_before = Pt(18)
                para.paragraph_format.space_after = Pt(12)
            elif is_clausula:
                run.bold = True
                run.font.size = Pt(fonte_tamanho)
                para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                para.paragraph_format.space_before = Pt(12)
                para.paragraph_format.space_after = Pt(6)
            elif is_assinatura:
                run.font.size = Pt(fonte_tamanho)
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.space_before = Pt(24)
            else:
                run.font.size = Pt(fonte_tamanho)
                para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                para.paragraph_format.space_after = Pt(espacamento.get("depois_paragrafo", 6))
            
            # Espaçamento entre linhas
            linha_spacing = espacamento.get("linha", 1.5)
            if isinstance(linha_spacing, (int, float)) and linha_spacing > 0:
                para.paragraph_format.line_spacing = linha_spacing
    
    def _is_titulo(self, texto: str) -> bool:
        """Verifica se é um título"""
        titulos = [
            'ALTERAÇÃO DO CONTRATO', 'CONTRATO SOCIAL', 'CONSOLIDAÇÃO',
            'ENCERRAMENTO', 'PREÂMBULO', 'QUALIFICAÇÃO DOS SÓCIOS',
            'QUADRO SOCIETÁRIO', 'CLÁUSULAS DE ALTERAÇÃO'
        ]
        texto_upper = texto.upper()
        return any(t in texto_upper for t in titulos) and len(texto) < 80
    
    def _is_clausula(self, texto: str) -> bool:
        """Verifica se é uma cláusula"""
        return bool(re.match(
            r'^(CLÁUSULA|PRIMEIRA|SEGUNDA|TERCEIRA|QUARTA|QUINTA|SEXTA|SÉTIMA|OITAVA|NONA|DÉCIMA|ARTIGO|ART\.)',
            texto,
            re.IGNORECASE
        ))
    
    def _is_assinatura(self, texto: str) -> bool:
        """Verifica se é linha de assinatura"""
        return texto.startswith('_') or 'CPF:' in texto or 'Assinatura' in texto


# Instância global
template_manager_fiel = TemplateManagerFiel()


def gerar_minuta_word(conteudo: str, dados_extraidos: Dict = None, template_path: str = None) -> bytes:
    """Função auxiliar para gerar documento Word fiel ao template"""
    return template_manager_fiel.gerar_documento_fiel(
        conteudo=conteudo,
        template_path=template_path,
        dados_extraidos=dados_extraidos
    )
